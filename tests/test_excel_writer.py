"""Tests for src/excel_writer.py: append-only safety, backups, flushing, lock errors."""
from __future__ import annotations

import datetime as dt
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from src.excel_writer import ExcelWriter, ExcelWriterError  # noqa: E402
from tests.helpers import make_workbook  # noqa: E402


class ExcelWriterTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.wb_path = os.path.join(self.tmp.name, "template.xlsx")
        make_workbook(self.wb_path)

    def _writer(self, flush_every_n: int = 1):
        return ExcelWriter(self.wb_path, "details", "Details_LineItems",
                           "_DocParser_ProcessedIDs",
                           os.path.join(self.tmp.name, "backups"),
                           flush_every_n=flush_every_n)

    def test_append_is_at_first_free_row_and_header_preserved(self):
        w = self._writer()
        idx = w.append_header_mapped_row({"FileID": "f1", "InvoiceNo": "INV-1",
                                          "NoSuchHeader": "ignored"})
        self.assertEqual(idx, 2)
        w.save()
        wb = load_workbook(self.wb_path)
        ws = wb["details"]
        self.assertEqual(ws.cell(row=1, column=1).value, "FileID")
        self.assertEqual(ws.cell(row=2, column=1).value, "f1")
        self.assertEqual(ws.cell(row=2, column=4).value, "INV-1")
        self.assertIsNone(ws.cell(row=2, column=2).value)  # FileName left blank

    def test_multiple_appends_sequence_rows(self):
        w = self._writer()
        a = w.append_header_mapped_row({"InvoiceNo": "INV-A"})
        b = w.append_header_mapped_row({"InvoiceNo": "INV-B"})
        self.assertEqual((a, b), (2, 3))
        w.save()

    def test_dates_written_as_real_types(self):
        w = self._writer()
        w.append_header_mapped_row({"InvoiceDate": "2026-09-15", "ProcessedAt": "2026-09-21T10:00:00"})
        w.save()
        wb = load_workbook(self.wb_path)
        ws = wb["details"]
        self.assertIsInstance(ws.cell(row=2, column=5).value, dt.date)
        self.assertIsInstance(ws.cell(row=2, column=3).value, dt.datetime)

    def test_backup_created_once(self):
        w = self._writer()
        w.append_header_mapped_row({"FileID": "f1"})
        w.save()
        w.save()  # second save: no new backup
        backups = os.listdir(os.path.join(self.tmp.name, "backups"))
        self.assertEqual(len(backups), 1)
        self.assertIn("template.xlsx", backups[0])

    def test_backup_is_a_copy_of_original(self):
        w = self._writer()
        w.append_header_mapped_row({"FileID": "f1"})
        w.save()
        bk = os.listdir(os.path.join(self.tmp.name, "backups"))[0]
        wb = load_workbook(os.path.join(self.tmp.name, "backups", bk))
        self.assertEqual(wb["details"].max_row, 1)  # backup == pristine template

    def test_line_items_written_under_first_headers(self):
        w = self._writer()
        rows = [["INV-1", "Vendor", "2026-09-15", "f1", "Item", 2, "Pcs",
                 5000.0, 10000.0, 900.0, 900.0, 11800.0, 1]]
        self.assertEqual(w.append_line_items(rows), 1)
        w.save()
        wb = load_workbook(self.wb_path)
        ws = wb["Details_LineItems"]
        self.assertEqual(ws.cell(row=2, column=1).value, "INV-1")
        self.assertEqual(ws.cell(row=2, column=5).value, "Item")

    def test_processed_id_appended(self):
        w = self._writer()
        w.append_processed_id("FID9")
        w.save()
        wb = load_workbook(self.wb_path)
        ws = wb["_DocParser_ProcessedIDs"]
        self.assertEqual(ws.cell(row=2, column=1).value, "FID9")

    def test_flush_every_n_batches(self):
        w = self._writer(flush_every_n=3)
        w.append_header_mapped_row({"InvoiceNo": "INV-1"})
        w.append_header_mapped_row({"InvoiceNo": "INV-2"})
        # no save should have hit disk before the threshold
        self.assertTrue(w._dirty)
        w.maybe_flush()  # 2 staged < 3 -> still not flushed
        self.assertTrue(w._dirty)
        self.assertFalse(os.path.exists(self.wb_path + ".tmp"))
        w.append_header_mapped_row({"InvoiceNo": "INV-3"})
        w.maybe_flush()  # 3 staged -> flushed
        self.assertFalse(w._dirty)
        wb = load_workbook(self.wb_path)
        self.assertEqual(wb["details"].max_row, 4)

    def test_save_wraps_permission_error(self):
        w = self._writer()
        w.append_header_mapped_row({"FileID": "f1"})

        def boom(path):
            raise PermissionError(13, "file in use by another process", w.path)

        orig = w.wb.save
        w.wb.save = boom
        try:
            with self.assertRaises(ExcelWriterError):
                w.save()
        finally:
            w.wb.save = orig

    def test_missing_worksheet_rejected(self):
        with self.assertRaises(ExcelWriterError):
            ExcelWriter(self.wb_path, "NoSuchSheet", "Details_LineItems",
                        "_DocParser_ProcessedIDs", self.tmp.name)

    def test_close_flushes_dirty_rows(self):
        w = self._writer()
        w.append_header_mapped_row({"InvoiceNo": "INV-1"})
        w.close()
        wb = load_workbook(self.wb_path)
        self.assertEqual(wb["details"].cell(row=2, column=4).value, "INV-1")


if __name__ == "__main__":
    unittest.main()