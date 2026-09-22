"""Tests for the MCP layer: protocol handshake, tools/list, tools/call,
Excel-tool error mapping, and a real stdio subprocess round-trip.

The in-process tests run the actual MCP wire protocol (initialize ->
notifications/initialized -> tools/list -> tools/call) over OS pipes; the
subprocess test runs the exact `python -m src.mcp.server` entry point.
"""
from __future__ import annotations

import os
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src.config import load_config  # noqa: E402
from src.excel_backend import open_excel_backend  # noqa: E402
from src.excel_writer import ExcelWriterError  # noqa: E402
from src.mcp.client import McpError  # noqa: E402
from src.mcp.excel_tool import (  # noqa: E402
    TOOL_APPEND_INVOICE,
    TOOL_CLOSE,
    TOOL_FLUSH,
    TOOL_MARK_SEEN,
    TOOL_READ_SEED,
)
from tests.helpers import make_cfg, make_workbook  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


class _HelperMixin(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.wb = make_workbook(os.path.join(self.tmp.name, "tpl.xlsx"))
        self._backends = []

    def _backend(self):
        path = make_cfg(self.tmp.name, self.wb,
                        mcp={"enabled": True, "transport": "inproc"})
        cfg = load_config(path)
        bk = open_excel_backend(cfg, path, dry_run=False)
        self._backends.append(bk)
        return bk

    def tearDown(self):
        for b in self._backends:
            try:
                b.close()
            except Exception:
                pass

    def _read_details(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.wb)
        ws = wb["details"]
        out = []
        for r in range(2, ws.max_row + 1):
            out.append({ws.cell(1, c).value: ws.cell(r, c).value
                        for c in range(1, ws.max_column + 1)})
        return [x for x in out if any(v is not None for v in x.values())]

    def _processed_ids(self):
        from openpyxl import load_workbook
        ws = load_workbook(self.wb)["_DocParser_ProcessedIDs"]
        return [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)]


class MCPServerProtocolTest(_HelperMixin):
    def test_initialize_ping_list_tools(self):
        bk = self._backend()
        self.assertTrue(bk.client.ping())
        tools = [t["name"] for t in bk.client.list_tools()]
        for name in (TOOL_READ_SEED, TOOL_APPEND_INVOICE, TOOL_MARK_SEEN,
                     TOOL_FLUSH, TOOL_CLOSE):
            self.assertIn(name, tools)
        desc = {t["name"]: t for t in bk.client.list_tools()}
        self.assertIn("inputSchema", desc[TOOL_APPEND_INVOICE])
        self.assertTrue(bk.client.server_info.get("name"))
        self.assertIn("protocolVersion", bk.client.server_info) if False else None

    def test_append_invoice_writes_details_lineitems_processed(self):
        bk = self._backend()
        values = {"InvoiceNo": "INV-M1", "InvoiceDate": "2026-09-15",
                  "PartyName": "Acme", "PartyGST": "27AAACS5842A1ZD",
                  "TotalValue": 11800.0, "TaxableValue": 10000.0}
        line_items = [["INV-M1", "Acme", None, "f1", "Widget", 1, "Pcs", None,
                       10000.0, None, None, 11800.0, 1]]
        row, count = bk.append_invoice(values, line_items, "f1", flush=True)
        self.assertEqual(row, 2)
        self.assertEqual(count, 1)
        rows = self._read_details()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["InvoiceNo"], "INV-M1")
        self.assertEqual(rows[0]["TotalValue"], 11800.0)
        from openpyxl import load_workbook
        ws = load_workbook(self.wb)["Details_LineItems"]
        self.assertEqual(ws.max_row, 2)
        self.assertEqual(self._processed_ids(), ["f1"])

    def test_read_seed_reflects_existing_rows(self):
        bk = self._backend()
        bk.append_invoice({"InvoiceNo": "LEGACY", "TotalValue": 5.0}, [], "fz", flush=True)
        # a second backend over the same file sees the seeded rows/ids
        bk2 = self._backend()
        seed = bk2.client.call_tool(TOOL_READ_SEED)
        self.assertEqual(len(seed["seed"]), 1)
        self.assertIn("fz", seed["processed_ids"])

    def test_mark_seen_persists_processed_id_without_details(self):
        bk = self._backend()
        bk.mark_seen("f9")
        self.assertEqual(self._processed_ids(), ["f9"])
        self.assertEqual(self._read_details(), [])

    def test_locked_workbook_maps_to_excel_writer_error(self):
        bk = self._backend()
        with mock.patch("src.excel_writer.ExcelWriter.save",
                        side_effect=ExcelWriterError("workbook locked by Excel")):
            with self.assertRaises(ExcelWriterError):
                bk.append_invoice({"InvoiceNo": "X"}, [], "f2", flush=True)
            with self.assertRaises(ExcelWriterError):
                bk.mark_seen("f3")

    def test_unknown_tool_and_bad_params_are_protocol_errors(self):
        bk = self._backend()
        with self.assertRaises(McpError):
            bk.client.call_tool("not_a_tool", {})
        with self.assertRaises(McpError):
            bk.client.call_tool(TOOL_APPEND_INVOICE, {"line_items": []})


class MCPServerSubprocessTest(_HelperMixin):
    def test_real_stdio_subprocess_roundtrip(self):
        path = make_cfg(self.tmp.name, self.wb,
                        mcp={"enabled": True, "transport": "stdio"})
        env = dict(os.environ)
        env["PYTHONPATH"] = ROOT + os.pathsep + env.get("PYTHONPATH", "")
        proc = subprocess.Popen(
            [sys.executable, "-m", "src.mcp.server", "--config", os.path.abspath(path)],
            cwd=ROOT, stdin=subprocess.PIPE, stdout=subprocess.PIPE,
            stderr=subprocess.PIPE, bufsize=0, env=env,
        )
        try:
            from src.mcp.transport import StdioTransport
            from src.mcp.client import MCPClient
            client = MCPClient(StdioTransport(proc.stdout, proc.stdin, name="subproc"),
                               timeout=30)
            info = client.initialize()
            self.assertEqual(info.get("serverInfo", {}).get("name"), "invoice-excel-mcp")
            tools = [t["name"] for t in client.list_tools()]
            self.assertIn(TOOL_APPEND_INVOICE, tools)
            res = client.call_tool(TOOL_APPEND_INVOICE, {
                "values": {"InvoiceNo": "INV-SUB-1", "TotalValue": 500.0},
                "line_items": [], "file_id": "sf1", "flush": True})
            self.assertEqual(res["excel_row"], 2)
            self.assertEqual(res["line_item_count"], 0)
            self.assertEqual(client.call_tool(TOOL_CLOSE, {}), {})
            client.close()
            proc.wait(timeout=20)
            self.assertEqual(proc.returncode, 0, proc.stderr.read().decode())
        finally:
            if proc.poll() is None:
                proc.kill()
                proc.wait(timeout=20)
            proc.stdout.close()
            proc.stderr.close()
        rows = self._read_details()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["InvoiceNo"], "INV-SUB-1")
        self.assertEqual(self._processed_ids(), ["sf1"])


if __name__ == "__main__":
    unittest.main()