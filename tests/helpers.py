"""Shared helpers for the audit test-suite.

* build_text_pdf()          -> a real one-page text PDF (readable by pypdf).
* FakeLMStudioServer        -> an in-process OpenAI-compatible server that
                               emulates LM Studio's /v1/chat/completions so the
                               real Extractor can be exercised without a model.
* make_workbook()/make_cfg() -> minimal Excel template + pipeline config.
"""
from __future__ import annotations

import json
import os
import posixpath
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse


# ---------------------------------------------------------------------------
# Minimal real text-PDF generator (no external libraries)
# ---------------------------------------------------------------------------
def _esc_pdf_text(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def build_text_pdf(text: str, path: str) -> str:
    objs = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj",
        (b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj"),
    ]
    stream = ("BT /F1 12 Tf 72 720 Td (%s) Tj ET" % _esc_pdf_text(text)).encode("latin-1", "replace")
    objs.append(b"4 0 obj << /Length %d >> stream\n%s\nendstream endobj" % (len(stream), stream))
    objs.append(b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj")

    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for obj in objs:
        offsets.append(len(data))
        data += obj + b"\n"
    xref_pos = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    for off in offsets[1:]:
        data += ("%010d 00000 n \n" % off).encode()
    data += b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n"
    data += str(xref_pos).encode() + b"\n%%EOF"
    with open(path, "wb") as fh:
        fh.write(bytes(data))
    return path


# ---------------------------------------------------------------------------
# In-process LM Studio emulator
# ---------------------------------------------------------------------------
class _Handler(BaseHTTPRequestHandler):
    server_version = "FakeLMStudio/1.0"
    protocol_version = "HTTP/1.1"

    def log_message(self, *args):  # silence
        pass

    def _read_body(self):
        length = int(self.headers.get("Content-Length", 0) or 0)
        return self.rfile.read(length)

    def do_GET(self):
        if urlparse(self.path).path == "/v1/models":
            self._json(200, {"object": "list", "data": [{"id": "qwen-test", "object": "model"}]})
        else:
            self._json(404, {"error": {"message": "not found", "type": "invalid_request_error"}})

    def do_POST(self):
        if urlparse(self.path).path != "/v1/chat/completions":
            self._json(404, {"error": {"message": "not found", "type": "invalid_request_error"}})
            return
        try:
            raw = self._read_body()
            body = json.loads(raw.decode("utf-8") or "{}")
        except Exception:
            self._json(400, {"error": {"message": "bad json", "type": "invalid_request_error"}})
            return
        server: FakeLMStudioServer = self.server  # type: ignore[attr-defined]
        server.requests.append(body)

        if server.mode == "no_response_format":
            if body.get("response_format"):
                self._json(400, {"error": {
                    "message": "response_format is not supported for this model",
                    "type": "invalid_request_error",
                    "param": "response_format"}})
                return
        elif server.mode == "fail_then_ok":
            if server.fail_count < server.fail_until:
                server.fail_count += 1
                self._json(500, {"error": {"message": "internal error", "type": "internal_error"}})
                return
        elif server.mode == "bad_json":
            server.payload_callback = lambda req: "this is definitely not JSON"

        content = server.payload_callback(body)
        if not isinstance(content, str):
            content = json.dumps(content, ensure_ascii=False)
        messages = body.get("messages", [])
        saw_image = any(
            isinstance(p, dict) and p.get("type") == "image_url" and "data:" in str(p.get("image_url", {}))
            for m in messages
            for p in (m.get("content", []) if isinstance(m.get("content"), list) else [])
        )
        payload = {
            "id": "chatcmpl-1234",
            "object": "chat.completion",
            "created": 0,
            "model": body.get("model", "qwen-test"),
            "system_fingerprint": "fp_test",
            "choices": [{
                "index": 0,
                "message": {"role": "assistant", "content": content},
                "finish_reason": "stop",
            }],
            "usage": {
                "prompt_tokens": 12,
                "completion_tokens": 7,
                "total_tokens": 19,
            },
        }
        if saw_image:
            payload["_saw_image"] = True
        self._json(200, payload)

    def _json(self, code: int, payload: Dict[str, Any]) -> None:
        data = json.dumps(payload).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


class FakeLMStudioServer:
    """Simple OpenAI-compatible chat-completions server used by extractor tests."""

    MODES = ("ok", "bad_json", "no_response_format", "fail_then_ok")

    def __init__(self, mode: str = "ok", payload=None, fail_until: int = 2):
        assert mode in self.MODES, mode
        self.mode = mode
        self.fail_until = fail_until
        self.fail_count = 0
        self.requests: List[Dict[str, Any]] = []
        default = {
            "invoice_number": "INV-STUB-1",
            "invoice_date": "2026-09-01",
            "vendor_name": "Stub Traders",
            "vendor_gstin": "27AAACS5842A1ZD",
            "taxable_value": 10000.0,
            "cgst_amount": 900.0,
            "sgst_amount": 900.0,
            "total_amount": 11800.0,
            "confidence": 0.98,
            "uncertain_fields": [],
            "line_items": [{"sl_no": 1, "description": "Item A", "qty": 1,
                            "unit": "Pcs", "taxable_value": 10000.0,
                            "cgst_amount": 900.0, "sgst_amount": 900.0,
                            "line_total": 11800.0}],
        }
        self.payload_callback = (lambda req: payload) if payload else (lambda req: default)
        self._httpd = ThreadingHTTPServer(("127.0.0.1", 0), _Handler)
        self._httpd.requests = self.requests
        self._httpd.mode = self.mode
        self._httpd.fail_count = 0
        self._httpd.fail_until = self.fail_until
        self._httpd.payload_callback = self.payload_callback
        self.port = self._httpd.server_address[1]
        self.base_url = f"http://127.0.0.1:{self.port}/v1"
        self._thread = threading.Thread(target=self._httpd.serve_forever, daemon=True)
        self._thread.start()

    def close(self) -> None:
        self._httpd.shutdown()
        self._httpd.server_close()

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        self.close()


# ---------------------------------------------------------------------------
# Excel template + config fixtures
# ---------------------------------------------------------------------------
def make_workbook(path: str, include_lineitems: bool = True,
                  include_processed: bool = True) -> str:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "details"
    ws.append(["FileID", "FileName", "ProcessedAt", "InvoiceNo", "InvoiceDate",
               "PartyName", "PartyGST", "VendorGST", "TotalValue", "TaxableValue",
               "CGSTAmt", "SGSTAmt", "_Confidence", "_ContentHash", "_SourceType"])
    if include_lineitems:
        ws2 = wb.create_sheet("Details_LineItems")
        ws2.append(["_InvoiceNo", "_PartyName", "_InvoiceDate", "_FileID",
                    "Description", "Qty", "Unit", "UnitRate", "TaxableValue",
                    "CGSTAmt", "SGSTAmt", "LineTotal", "SlNo"])
    if include_processed:
        ws3 = wb.create_sheet("_DocParser_ProcessedIDs")
        ws3.append(["ID", "AddedAt"])
    wb.save(path)
    return path


def make_cfg(tmpdir: str, wb_path: str, folder_id: str = "fakefolder",
             **overrides: Any) -> str:
    """Write a pipeline config into tmpdir and return its path."""
    cfg: Dict[str, Any] = {
        "drive": {
            "service_account_json": os.path.abspath("config.json"),
            "folder_id": folder_id,
            "recursive": True,
            "allowed_extensions": [".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff"],
            "download_dir": os.path.join(tmpdir, "downloads"),
            "max_file_size_mb": 50,
            "delete_temp_files": False,
        },
        "excel": {
            "template_path": os.path.abspath(wb_path),
            "target_sheet": "details",
            "lineitems_sheet": "Details_LineItems",
            "processed_ids_sheet": "_DocParser_ProcessedIDs",
            "backup_dir": os.path.join(tmpdir, "backups"),
            "write_line_items": True,
            "write_processed_ids": True,
            "flush_every_n": 1,
        },
        "model": {
            "base_url": "http://localhost:1234/v1",
            "api_key": "lm-studio",
            "model": "qwen-test",
            "vision_enabled": True,
            "retries": 0,
            "timeout_seconds": 10,
        },
        "extraction": {"require_invoice_number_to_append": True},
        "duplicates": {
            "ledger_file": os.path.join(tmpdir, "ledger.jsonl"),
            "check_invoice_fingerprint": True,
        },
        "filtering": {
            "min_file_size_kb": 25,
            "enable_content_filter": True,
            "enable_logo_only_filter": True,
            "min_invoice_signal_score": 2,
        },
        "audit": {
            "dir": os.path.join(tmpdir, "audit"),
            "consolidated_csv": os.path.join(tmpdir, "audit", "audit_all_runs.csv"),
        },
        "column_map": {
            "file_id": "FileID",
            "file_name": "FileName",
            "processed_at": "ProcessedAt",
            "invoice_number": "InvoiceNo",
            "invoice_date": "InvoiceDate",
            "vendor_name": "PartyName",
            "vendor_gstin": ["PartyGST", "VendorGST"],
            "total_amount": "TotalValue",
            "taxable_value": "TaxableValue",
            "cgst_amount": "CGSTAmt",
            "sgst_amount": "SGSTAmt",
        },
    }
    cfg.update(overrides)
    path = os.path.join(tmpdir, "cfg.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(cfg, fh)
    return path


def dump(x) -> str:
    return "" if x is None else str(x)