"""Excel template reader/writer.

Guarantees:
  * never renames / deletes / reorders / adds worksheets;
  * never changes existing column names (writes only under existing headers);
  * preserves the workbook (saved with a safety backup once per run);
  * appends rows at the first free row below data.
"""
from __future__ import annotations

import datetime as dt
import logging
import os
import shutil
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

log = logging.getLogger("invoice_pipeline.excel")


class ExcelWriterError(Exception):
    pass


class ExcelWriter:
    def __init__(self, path: str, target_sheet: str, lineitems_sheet: str,
                 processed_ids_sheet: str, backup_dir: str, flush_every_n: int = 1):
        self.path = os.path.abspath(path)
        self.target_sheet = target_sheet
        self.lineitems_sheet = lineitems_sheet
        self.processed_ids_sheet = processed_ids_sheet
        self.backup_dir = backup_dir
        self.wb = load_workbook(self.path, data_only=False)
        self._backup_done = False
        self._lineitems_headers: List[str] = []
        self._processed_headers: List[str] = []
        self._dirty = False
        # how many staged rows before the workbook is flushed to disk
        # (1 = save after every invoice; raise e.g. to 10 for large batches)
        self.flush_every_n = max(1, int(flush_every_n or 1))
        self._staged_rows = 0

        for sheet in (target_sheet, lineitems_sheet, processed_ids_sheet):
            if sheet not in self.wb.sheetnames:
                raise ExcelWriterError(
                    f"Missing expected worksheet {sheet!r}. The template may have changed; adjust config."
                )

        self.ws = self.wb[target_sheet]
        self.headers = self._read_headers(self.ws)
        if not self.headers:
            raise ExcelWriterError(f"Target sheet {target_sheet!r} has no header row.")
        self.header_index = {name: i for i, name in enumerate(self.headers) if name}
        if self.lineitems_sheet in self.wb.sheetnames:
            self._lineitems_headers = self._read_headers(self.wb[self.lineitems_sheet])
        if self.processed_ids_sheet in self.wb.sheetnames:
            self._processed_headers = self._read_headers(self.wb[self.processed_ids_sheet])

    @property
    def line_items_headers(self) -> List[str]:
        """Ordered header names of the line-items worksheet."""
        return list(self._lineitems_headers)

    @property
    def processed_ids_headers(self) -> List[str]:
        """Ordered header names of the processed-IDs worksheet."""
        return list(self._processed_headers)

    # ------------------------------------------------------------------
    @staticmethod
    def _read_headers(ws) -> List[str]:
        row = ws[1]
        vals = []
        for cell in row:
            vals.append(str(cell.value).strip() if cell.value is not None else "")
        # trim trailing empties
        while vals and vals[-1] == "":
            vals.pop()
        return vals

    def first_free_row(self, ws=None) -> int:
        ws = ws or self.ws
        return ws.max_row + 1

    # ------------------------------------------------------------------
    def append_header_mapped_row(self, values: Dict[str, Any]) -> int:
        """values: {column header name: value}. Only known headers are written."""
        row_idx = self.first_free_row()
        for col_name, value in values.items():
            if col_name in self.header_index:
                cell = self.ws.cell(row=row_idx, column=self.header_index[col_name] + 1)
                cell.value = self._convert_for_cell(col_name, value)
                self._apply_number_format(cell, col_name)
        self._dirty = True
        self._staged_rows += 1
        log.info("details: row %d staged for %s", row_idx, values.get("InvoiceNo") or values.get("FileID"))
        return row_idx

    def append_line_items(self, rows: List[List[Any]]) -> int:
        if not self.lineitems_sheet or not rows:
            return 0
        ws = self.wb[self.lineitems_sheet]
        ncols = len(self._lineitems_headers)
        start = self.first_free_row(ws)
        for r, row in enumerate(rows):
            for c in range(ncols):
                if c < len(row) and row[c] is not None:
                    ws.cell(row=start + r, column=c + 1).value = row[c]
        self._dirty = True
        self._staged_rows += len(rows)
        return len(rows)

    def append_processed_id(self, file_id: str) -> int:
        if not self.processed_ids_sheet:
            return 0
        ws = self.wb[self.processed_ids_sheet]
        row_idx = self.first_free_row(ws)
        ws.cell(row=row_idx, column=1).value = file_id
        ws.cell(row=row_idx, column=2).value = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self._dirty = True
        self._staged_rows += 1
        return row_idx

    # ------------------------------------------------------------------
    def maybe_flush(self) -> str:
        """Persist to disk once 'flush_every_n' rows have been staged."""
        if self._dirty and self._staged_rows >= self.flush_every_n:
            return self.save()
        return self.path

    # ------------------------------------------------------------------
    @staticmethod
    def _date_cols() -> set:
        return {"InvoiceDate", "DueDate", "AckDate", "_EmailDate"}

    def _convert_for_cell(self, col_name: str, value: Any) -> Any:
        if value is None:
            return None
        if isinstance(value, dt.datetime):
            return value
        if col_name in ("ProcessedAt", "_ProcessedAt"):
            if isinstance(value, str):
                try:
                    return dt.datetime.fromisoformat(value)
                except ValueError:
                    return value
            return value
        if col_name in self._date_cols():
            if isinstance(value, str) and len(value) == 10 and value[4] == "-":
                try:
                    return dt.date.fromisoformat(value)
                except ValueError:
                    return value
            return value
        if isinstance(value, (dt.date,)):
            return value
        # drop stray leading zero strings that Excel would store as text-only? keep simple
        return value

    @staticmethod
    def _apply_number_format(cell, col_name: str) -> None:
        if col_name in ("ProcessedAt", "_ProcessedAt"):
            cell.number_format = "YYYY-MM-DD HH:MM:SS"
        elif col_name in ("InvoiceDate", "DueDate", "AckDate"):
            cell.number_format = "DD-MM-YYYY"

    # ------------------------------------------------------------------
    def backup_once(self) -> Optional[str]:
        if self._backup_done:
            return None
        os.makedirs(self.backup_dir, exist_ok=True)
        ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = os.path.join(self.backup_dir, f"backup_{ts}_{os.path.basename(self.path)}")
        shutil.copy2(self.path, dest)
        self._backup_done = True
        return dest

    def save(self) -> str:
        if not self._dirty:
            return self.path
        self.backup_once()
        tmp = self.path + ".tmp"
        try:
            self.wb.save(tmp)
            os.replace(tmp, self.path)
        except OSError as exc:
            # e.g. the workbook is open in Excel on Windows (PermissionError)
            raise ExcelWriterError(
                f"Could not save workbook {self.path}: {exc}. "
                "Close the file in Excel if it is open, then re-run."
            ) from exc
        self._dirty = False
        self._staged_rows = 0
        log.info("Excel saved: %s", self.path)
        return self.path

    def close(self) -> None:
        if self._dirty:
            self.save()

    # ------------------------------------------------------------------
    def read_template_for_seed(self) -> List[Dict[str, Any]]:
        """Materialize existing rows as dicts for duplicate-checker seeding."""
        def idx(name):  # locate by header name
            return self.header_index.get(name)
        out: List[Dict[str, Any]] = []
        for r in range(2, self.ws.max_row + 1):
            row: Dict[str, Any] = {"row": r}
            def val(name):
                i = idx(name)
                return self.ws.cell(row=r, column=i + 1).value if i is not None else None
            row["file_id"] = val("FileID") or val("_FileID") or None
            row["content_hash"] = val("_ContentHash") or None
            row["invoice_number"] = val("InvoiceNo") or None
            row["vendor_gstin"] = val("VendorGST") or val("PartyGST") or None
            d = val("InvoiceDate")
            row["invoice_date"] = d.strftime("%Y-%m-%d") if isinstance(d, (dt.date, dt.datetime)) else d
            row["total_amount"] = val("TotalValue") or None
            fp = _fingerprint_from_parts(row["invoice_number"], row["vendor_gstin"],
                                         row["invoice_date"], row["total_amount"])
            row["fingerprint"] = fp
            if row["file_id"] or row["content_hash"] or fp:
                out.append(row)
        return out

    def count_data_rows(self) -> int:
        return max(0, self.ws.max_row - 1) if self.ws.max_row else 0


def _fingerprint_from_parts(invoice_no, gstin, date, total) -> Optional[str]:
    parts = []
    parts.append(str(invoice_no or "").strip().upper().replace(" ", ""))
    parts.append(str(gstin or "").strip().upper().replace(" ", ""))
    parts.append(str(date or "").strip())
    if total is not None:
        try:
            parts.append(f"{round(float(total), 2):.2f}")
        except (TypeError, ValueError):
            parts.append("")
    else:
        parts.append("")
    return "|".join(parts) if any(parts) else None